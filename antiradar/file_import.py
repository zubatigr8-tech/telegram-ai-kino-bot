"""Tayyor ro'yxatdan (Excel .xlsx yoki .csv) radar, kamera va yo'l belgilarini import qilish.

Admin botga fayl yuboradi. Rasmiy fotoradarlar ro'yxati (Viloyat / Tuman / Joylashgan joyi /
Turi / Kenglik / Uzunlik) to'g'ridan-to'g'ri o'qiladi. Sarlavhalar o'zbek, rus yoki ingliz tilida,
"Kenglik / Широта / Latitude" kabi aralash ko'rinishda ham bo'lishi mumkin.

Ustunlar (tartib ahamiyatsiz):
    lat, lon       — koordinatalar (majburiy)
    kind / turi    — "Statsionar radar", "Statsionar kamera", "Intellektual + radar"
                     yoki ichki nom (fixed, camera, crossing...). Standart: fixed
    speed_limit    — tezlik chegarasi, km/soat (ixtiyoriy)
    direction      — kamera o'lchaydigan yo'nalish, 0-360° (ixtiyoriy)
Bir xil turdagi nuqta 15 m ichida allaqachon bo'lsa, takror qo'shilmaydi.
"""
import csv
import io
import re
from dataclasses import dataclass

from sqlalchemy import select, update

from antiradar.alerts import RADAR_KINDS, SIGN_KINDS
from antiradar.config import settings
from antiradar.db import Camera, get_session
from antiradar.geo import distance_m

DUPLICATE_RADIUS_M = 15
HEADER_SEARCH_ROWS = 15
FILE_SOURCE = "file"
ALL_KINDS = set(RADAR_KINDS) | set(SIGN_KINDS)
HEADER_ALIASES = {
    "lat": "lat", "latitude": "lat", "широта": "lat", "kenglik": "lat",
    "lon": "lon", "lng": "lon", "longitude": "lon", "долгота": "lon", "uzunlik": "lon",
    "kind": "kind", "type": "kind", "a type": "kind", "тип": "kind", "turi": "kind",
    "speed_limit": "speed_limit", "limit": "speed_limit", "maxspeed": "speed_limit",
    "ограничение": "speed_limit", "tezlik": "speed_limit",
    "direction": "direction", "направление": "direction", "yonalish": "direction",
}


@dataclass
class ParsedPoint:
    lat: float
    lon: float
    kind: str
    speed_limit: int | None
    direction: float | None


def _column_key(header) -> str | None:
    """'Kenglik / Широта / Latitude' → 'lat'. Har bir '/' bo'lagi alohida tekshiriladi."""
    for part in str(header or "").split("/"):
        key = HEADER_ALIASES.get(re.sub(r"\s+", " ", part).strip().lower().replace("'", "").replace("ʻ", ""))
        if key:
            return key
    return None


def map_kind(raw: str | None) -> str | None:
    """Rasmiy ro'yxatdagi tur nomini ichki turga aylantiradi. None — tanilmadi."""
    value = (raw or "").strip().lower()
    if not value:
        return "fixed"
    if value in ALL_KINDS:
        return value
    if "intellekt" in value or "интеллект" in value or "smart" in value:
        return "smart"
    if "radar" in value or "радар" in value:
        return "fixed"
    if "kamera" in value or "камера" in value or "camera" in value:
        return "camera"
    return None


def _num(raw) -> float | None:
    if isinstance(raw, (int, float)):
        return float(raw)
    raw = str(raw or "").strip().replace(",", ".")
    return float(raw) if raw else None


def parse_rows(rows: list[list]) -> tuple[list[ParsedPoint], list[str]]:
    """Jadval qatorlaridan (sarlavha qayerda bo'lsa ham) nuqtalarni ajratadi. (nuqtalar, xatolar)."""
    header_index, columns = None, {}
    for i, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        keys = {idx: _column_key(cell) for idx, cell in enumerate(row)}
        keys = {idx: key for idx, key in keys.items() if key}
        if "lat" in keys.values() and "lon" in keys.values():
            header_index, columns = i, keys
            break
    if header_index is None:
        return [], ["Faylda kenglik (lat) va uzunlik (lon) ustunlari topilmadi"]

    points, errors = [], []
    for line_no, raw_row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        row = {key: raw_row[idx] for idx, key in columns.items() if idx < len(raw_row)}
        if not any(ch.isdigit() for ch in str(row.get("lat") or "")):
            continue  # bo'sh yoki ikkinchi sarlavha qatori ("Kenglik / Широта / Latitude")
        try:
            lat, lon = _num(row.get("lat")), _num(row.get("lon"))
            if lat is None or lon is None or not (-90 <= lat <= 90 and -180 <= lon <= 180):
                raise ValueError("koordinata noto'g'ri")
            kind = map_kind(row.get("kind"))
            if kind is None:
                raise ValueError(f"noma'lum tur: {row.get('kind')}")
            limit = _num(row.get("speed_limit"))
            direction = _num(row.get("direction"))
            points.append(
                ParsedPoint(
                    lat=lat,
                    lon=lon,
                    kind=kind,
                    speed_limit=int(limit) if limit else None,
                    direction=direction % 360 if direction is not None else None,
                )
            )
        except ValueError as exc:
            errors.append(f"{line_no}-qator: {exc}")
    return points, errors


def parse_csv(text: str) -> tuple[list[ParsedPoint], list[str]]:
    text = text.lstrip("﻿")
    first_line = text.splitlines()[0] if text else ""
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    return parse_rows(list(csv.reader(io.StringIO(text), delimiter=delimiter)))


def parse_xlsx(data: bytes) -> tuple[list[ParsedPoint], list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    points, errors = [], []
    for sheet in workbook.worksheets:
        sheet_points, sheet_errors = parse_rows([list(r) for r in sheet.iter_rows(values_only=True)])
        points += sheet_points
        if sheet_points or len(workbook.worksheets) == 1:
            errors += sheet_errors
    return points, errors


def parse_file(file_name: str, data: bytes) -> tuple[list[ParsedPoint], list[str]]:
    if file_name.lower().endswith(".xlsx"):
        return parse_xlsx(data)
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("cp1251")  # Excel'ning ruscha Windows kodirovkasi
    return parse_csv(text)


async def clear_imported() -> int:
    """Fayldan yuklangan barcha nuqtalarni o'chiradi (yangi ro'yxatni to'liq qayta yuklashdan oldin)."""
    async with get_session() as session:
        result = await session.execute(
            update(Camera).where(Camera.source == FILE_SOURCE, Camera.is_active.is_(True)).values(is_active=False)
        )
        await session.commit()
        return result.rowcount or 0


async def save_points(points: list[ParsedPoint]) -> tuple[int, int]:
    """(qo'shildi, takror sababli o'tkazib yuborildi) qaytaradi."""
    added = skipped = 0
    async with get_session() as session:
        existing = [
            (c.lat, c.lon, c.kind)
            for c in await session.scalars(select(Camera).where(Camera.is_active.is_(True)))
        ]
        for p in points:
            if any(
                kind == p.kind and abs(lat - p.lat) < 0.001 and distance_m(lat, lon, p.lat, p.lon) <= DUPLICATE_RADIUS_M
                for lat, lon, kind in existing
            ):
                skipped += 1
                continue
            session.add(
                Camera(
                    country=settings.COUNTRY,
                    lat=p.lat,
                    lon=p.lon,
                    kind=p.kind,
                    speed_limit=p.speed_limit,
                    direction=p.direction,
                    source=FILE_SOURCE,
                )
            )
            existing.append((p.lat, p.lon, p.kind))
            added += 1
        await session.commit()
    return added, skipped
